#!/usr/bin/env python3
"""
manuscript/edits/2026-09-03_simulation_procedure.py

Point 1 of the 2026-09-03 review pass: the description of the simulation procedure was
incomplete -- no time step, no numerical scheme, no language/tooling, and no statement
that the same detector used on the recorded EMG is applied to the simulated trace. This
inserts one new paragraph, sourced from scripts/delaymodel/sim.py (DT, the Euler-Maruyama
updates, the response criterion, the _bursts detector) and scripts/README.md ("Sampling
rate -- fixed 2026-08-30"), immediately after the paragraph that introduces simulation's
role in the paper ("Simulation enters this paper in one role...") and before the
"Evaluation over the parameter space" subsubsection heading.

Formatting note found while doing this: that anchor paragraph is Normal style, a single
run, color 1F4E9E (blue = new since R1_v3), with explicit <w:b w:val="0"/><w:i w:val="0"/>
-- NOT bold or italic, despite docx2tex.py rendering it as \textbf{\emph{...}} in the
LaTeX copy. That is a docx2tex.py bug (it checks `rPr.find(qn('w:b')) is not None`, which
is true even when w:val="0" explicitly turns bold off) -- reported separately, not fixed
here. The new paragraph matches the anchor's actual (plain, blue) formatting, not the
LaTeX copy's misleading rendering of it.

Usage: python3 add_sim_procedure.py <path-to-docx>   (edits in place)
"""
import sys
import docx
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from lxml import etree

ANCHOR = "no quantity reported below is conditional on a fitted value"

NEW_TEXT = (
    "The equations above were integrated numerically in Python (NumPy) by the explicit "
    "Euler–Maruyama scheme, at a fixed time step of 1/1024 s — the EMG and EEG "
    "acquisition rate (Methods, EMG acquisition and preprocessing) — so that the "
    "detection procedure below runs on the same time grid as the recorded signals. Each "
    "simulated trial yields one continuous trace of the gated drive, from stimulus onset "
    "to the earlier of the response or the trial's end; the response time is the first "
    "time step at which this drive reaches the response criterion, after which the fixed "
    "residual motor latency is added, as in the immediate-report architecture. Partial "
    "muscle activations are extracted from this trace by the same rule applied to the "
    "recorded EMG (Methods, EMG acquisition and preprocessing): the drive crossing above "
    "zero, merging of intervals separated by 25 ms or less, and discarding merged "
    "intervals with fewer than an adaptively set minimum of supra-threshold samples. The "
    "one part of that procedure not simulated is the integrated-profile boundary "
    "refinement, which relocates interval boundaries without changing their number, and "
    "only the number is used here. Trials are simulated independently per condition from "
    "a fixed pseudo-random seed, so every reported quantity is exactly reproducible."
)


def insert_paragraph_after(paragraph, text, style_name, bold, italic, color_hex):
    new_p_elm = OxmlElement("w:p")
    paragraph._p.addnext(new_p_elm)
    new_para = Paragraph(new_p_elm, paragraph._parent)
    new_para.style = style_name
    run = new_para.add_run(text)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = docx.shared.RGBColor.from_string(color_hex)
    return new_para


def main():
    path = sys.argv[1] if len(sys.argv) > 1 else \
        "Achard_et_al_R1_v3_tracked-changes_2026-09-02.docx"
    d = docx.Document(path)

    hits = [i for i, p in enumerate(d.paragraphs) if ANCHOR in p.text]
    if len(hits) != 1:
        raise SystemExit(f"expected exactly one anchor paragraph, found {len(hits)}: {hits}")
    idx = hits[0]
    anchor_p = d.paragraphs[idx]

    # sanity-check the anchor's own formatting so the new paragraph really matches it,
    # rather than trusting the docstring's claim
    r = anchor_p.runs[-1]
    rPr = r._r.find(qn("w:rPr"))
    b_el = rPr.find(qn("w:b"))
    i_el = rPr.find(qn("w:i"))
    anchor_bold = b_el is not None and b_el.get(qn("w:val")) not in ("0", "false")
    anchor_italic = i_el is not None and i_el.get(qn("w:val")) not in ("0", "false")
    anchor_color = r.font.color.rgb if r.font.color and r.font.color.type else None
    assert anchor_p.style.name == "Normal", anchor_p.style.name
    assert not anchor_bold and not anchor_italic, (anchor_bold, anchor_italic)
    assert str(anchor_color) == "1F4E9E", anchor_color

    next_p = d.paragraphs[idx + 1]
    assert next_p.style.name == "Heading 3", next_p.style.name
    assert "Evaluation over the parameter space" in next_p.text

    new_para = insert_paragraph_after(
        anchor_p, NEW_TEXT, style_name="Normal",
        bold=False, italic=False, color_hex="1F4E9E",
    )

    d.save(path)

    # re-open and verify
    d2 = docx.Document(path)
    hits2 = [i for i, p in enumerate(d2.paragraphs) if ANCHOR in p.text]
    assert len(hits2) == 1
    idx2 = hits2[0]
    inserted = d2.paragraphs[idx2 + 1]
    assert NEW_TEXT[:40] in inserted.text, inserted.text[:80]
    following = d2.paragraphs[idx2 + 2]
    assert following.style.name == "Heading 3" and "Evaluation over the parameter space" in following.text
    print("OK: inserted paragraph at index", idx2 + 1, "length", len(inserted.text))
    print("style:", inserted.style.name)
    r2 = inserted.runs[0]
    print("bold:", r2.bold, "italic:", r2.italic, "color:", r2.font.color.rgb)


if __name__ == "__main__":
    main()
